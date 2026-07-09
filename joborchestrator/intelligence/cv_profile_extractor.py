from __future__ import annotations

import json
import logging
import os
import re
from io import BytesIO
from typing import Any

import httpx
import yaml

NVIDIA_BASE_URL = os.getenv("NVIDIA_BASE_URL") or "https://integrate.api.nvidia.com/v1"
DEFAULT_PROFILE_EXTRACTION_MODEL = (
    os.getenv("PROFILE_EXTRACTION_MODEL")
    or os.getenv("NVIDIA_RANKING_MODEL")
    or os.getenv("NVIDIA_MODEL")
    or "nvidia/llama-3.3-nemotron-super-49b-v1"
)
DEFAULT_PROFILE_EXTRACTION_TIMEOUT_SECONDS = float(os.getenv("PROFILE_EXTRACTION_TIMEOUT_SECONDS", "180"))
DEFAULT_PROFILE_EXTRACTION_VALIDATION_RETRIES = int(os.getenv("PROFILE_EXTRACTION_VALIDATION_RETRIES", "1"))

PROFILE_SCHEMA_VERSION = 1
logger = logging.getLogger(__name__)


class CVProfileError(RuntimeError):
    pass


def extract_text_from_cv(filename: str, content: bytes) -> str:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix == "pdf":
        return _extract_pdf_text(content)
    if suffix == "docx":
        return _extract_docx_text(content)
    if suffix in {"txt", "md"}:
        return content.decode("utf-8", errors="ignore")
    raise CVProfileError("Upload a CV as PDF, DOCX, TXT, or MD.")


def build_profile_from_cv_text(
    cv_text: str,
    *,
    model: str = DEFAULT_PROFILE_EXTRACTION_MODEL,
    timeout: float = DEFAULT_PROFILE_EXTRACTION_TIMEOUT_SECONDS,
) -> dict[str, Any]:
    text = cv_text.strip()
    if len(text) < 200:
        raise CVProfileError("The CV text is too short to extract a useful profile.")
    payload = {
        "goal": (
            "Extract a career-agnostic job search profile from this CV. "
            "The result must work for programmers, designers, sales, operations, finance, healthcare, or any career."
        ),
        "rules": [
            "Infer likely target roles from evidence in the CV, not from stereotypes.",
            "Use a maximum of 8 top-level skill domains. Prefer stable domain names users can reuse.",
            "Assign skill.level using this rubric: strong = explicit usage in 2+ contexts or outcomes/metrics; medium = mentioned with one concrete context; weak = keyword/list item with no usage context.",
            "Do not invent employers, degrees, years, or seniority.",
            "Populate dealbreakers and preferred_locations only when explicitly stated. Return empty arrays when not stated.",
            "Include evidence for every skill and industry. Mark ambiguity in extraction_notes instead of resolving silently.",
            "Prefer concise labels users can edit later.",
            "Return only JSON.",
        ],
        "cv_text": text[:24000],
    }
    key = os.getenv("NVIDIA_API_KEY") or os.getenv("NIM_API_KEY")
    if not key:
        raise CVProfileError("NVIDIA_API_KEY or NIM_API_KEY is required to analyze CVs with AI.")
    validation_feedback: str | None = None
    for attempt in range(DEFAULT_PROFILE_EXTRACTION_VALIDATION_RETRIES + 1):
        raw = _call_nvidia_profile_extraction(payload, key, model, timeout, validation_feedback)
        profile = normalize_profile_payload(_extract_json_object(raw))
        validation_feedback = _profile_validation_error(profile)
        if not validation_feedback:
            return profile
        if attempt < DEFAULT_PROFILE_EXTRACTION_VALIDATION_RETRIES:
            logger.warning("Retrying CV profile extraction after invalid response: %s", validation_feedback)
            continue
        raise CVProfileError(f"AI profile extraction response was incomplete: {validation_feedback}")
    raise CVProfileError("AI profile extraction did not produce a usable profile.")


def _call_nvidia_profile_extraction(
    payload: dict[str, Any],
    api_key: str,
    model: str,
    timeout: float,
    validation_feedback: str | None = None,
) -> str:
    user_payload = {
        "input": payload,
        "output_shape": _profile_shape(),
    }
    if validation_feedback:
        user_payload["previous_response_error"] = validation_feedback
        user_payload["instruction"] = "Return a corrected complete JSON object only."
    try:
        response = httpx.post(
            f"{NVIDIA_BASE_URL.rstrip('/')}/chat/completions",
            headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
            json={
                "model": model,
                "temperature": 0.1,
                "response_format": {"type": "json_object"},
                "messages": [
                    {
                        "role": "system",
                        "content": (
                            "You are a strict career profile extraction engine. "
                            "Return only valid JSON matching the requested shape. "
                            "Do not return placeholders such as 'strong | medium | weak'."
                        ),
                    },
                    {"role": "user", "content": json.dumps(user_payload, ensure_ascii=False)},
                ],
            },
            timeout=timeout,
        )
        response.raise_for_status()
    except httpx.HTTPStatusError as exc:
        detail = exc.response.text[:1000] if exc.response is not None else ""
        raise CVProfileError(f"NVIDIA CV analysis failed: status={exc.response.status_code} body={detail!r}") from exc
    except httpx.HTTPError as exc:
        raise CVProfileError(f"NVIDIA CV analysis failed: {type(exc).__name__}: {exc!r}") from exc
    return str(response.json()["choices"][0]["message"]["content"])


def normalize_profile_payload(payload: dict[str, Any]) -> dict[str, Any]:
    skills = []
    for item in payload.get("skills") or []:
        if not isinstance(item, dict):
            continue
        name = str(item.get("name") or "").strip()
        if not name:
            continue
        level = str(item.get("level") or "medium").lower()
        if level not in {"strong", "medium", "weak"}:
            level = "medium"
        skills.append(
            {
                "name": name,
                "category": str(item.get("category") or "General").strip() or "General",
                "level": level,
                "evidence": str(item.get("evidence") or "").strip(),
            }
        )
    skills.sort(key=lambda item: (item["category"].lower(), {"strong": 0, "medium": 1, "weak": 2}[item["level"]], item["name"].lower()))

    return {
        "schema_version": PROFILE_SCHEMA_VERSION,
        "headline": str(payload.get("headline") or "").strip(),
        "target_roles": _clean_list(payload.get("target_roles")),
        "secondary_roles": _clean_list(payload.get("secondary_roles")),
        "role_aliases": _clean_role_aliases(payload.get("role_aliases")),
        "skills": skills,
        "industries": _clean_list(payload.get("industries")),
        "preferred_locations": _clean_list(payload.get("preferred_locations")),
        "preferred_work_modes": _clean_list(payload.get("preferred_work_modes")),
        "dealbreakers": _clean_list(payload.get("dealbreakers")),
        "avoid_roles": _clean_list(payload.get("avoid_roles")),
        "real_experience_years": _number(payload.get("real_experience_years"), 0.0),
        "notes": str(payload.get("notes") or "").strip(),
        "suggested_roles_reasoning": str(payload.get("suggested_roles_reasoning") or "").strip(),
        "extraction_notes": _clean_list(payload.get("extraction_notes")),
        "confidence": _confidence(payload.get("confidence")),
    }


def profile_payload_to_candidate_profile(profile: dict[str, Any]) -> dict[str, Any]:
    skills = profile.get("skills") or []
    return {
        "target_roles": _clean_list(profile.get("target_roles")),
        "secondary_roles": _clean_list(profile.get("secondary_roles")),
        "role_aliases": _clean_role_aliases(profile.get("role_aliases")),
        "strong_skills": [skill["name"] for skill in skills if skill.get("level") == "strong"],
        "medium_skills": [skill["name"] for skill in skills if skill.get("level") == "medium"],
        "weak_skills": [skill["name"] for skill in skills if skill.get("level") == "weak"],
        "industries": _clean_list(profile.get("industries")),
        "preferred_locations": _clean_list(profile.get("preferred_locations")),
        "preferred_work_modes": _clean_list(profile.get("preferred_work_modes")),
        "dealbreakers": _clean_list(profile.get("dealbreakers")),
        "avoid_roles": _clean_list(profile.get("avoid_roles")),
        "real_experience_years": _number(profile.get("real_experience_years"), 0.0),
        "notes": str(profile.get("notes") or profile.get("headline") or "").strip() or None,
    }


def _extract_pdf_text(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ModuleNotFoundError as exc:
        raise CVProfileError("PDF upload requires pypdf.") from exc
    reader = PdfReader(BytesIO(content))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx_text(content: bytes) -> str:
    try:
        from docx import Document
    except ModuleNotFoundError as exc:
        raise CVProfileError("DOCX upload requires python-docx.") from exc
    document = Document(BytesIO(content))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _extract_json_object(text: str) -> dict[str, Any]:
    raw = text.strip()
    if raw.startswith("```"):
        raw = re.sub(r"^```(?:json)?", "", raw).strip()
        raw = re.sub(r"```$", "", raw).strip()
    match = re.search(r"\{.*\}", raw, flags=re.DOTALL)
    if not match:
        raise CVProfileError("AI response did not include a JSON object.")
    try:
        parsed = _loads_lenient_json(match.group(0))
    except json.JSONDecodeError as exc:
        raise CVProfileError(f"AI response JSON was invalid: {exc}") from exc
    if not isinstance(parsed, dict):
        raise CVProfileError("AI response JSON must be an object.")
    return parsed


def _loads_lenient_json(raw: str) -> dict[str, Any]:
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        repaired = _repair_common_json_issues(raw)
    try:
        return json.loads(repaired)
    except json.JSONDecodeError:
        parsed = yaml.safe_load(repaired)
        if isinstance(parsed, dict):
            return parsed
        raise


def _repair_common_json_issues(raw: str) -> str:
    repaired = raw.strip()
    repaired = re.sub(r",(\s*[}\]])", r"\1", repaired)
    repaired = re.sub(r"(?<!\\)'([^'\\]*(?:\\.[^'\\]*)*)'", lambda m: json.dumps(m.group(1)), repaired)
    return repaired


def _clean_list(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    seen = set()
    cleaned = []
    for item in value:
        text = str(item or "").strip()
        key = text.lower()
        if text and key not in seen:
            cleaned.append(text)
            seen.add(key)
    return cleaned


def _clean_role_aliases(value: Any) -> dict[str, list[str]]:
    if not isinstance(value, dict):
        return {}
    cleaned: dict[str, list[str]] = {}
    for role, aliases in value.items():
        role_name = str(role or "").strip()
        values = _clean_list(aliases)
        if role_name and values:
            cleaned[role_name] = values
    return cleaned


def _number(value: Any, fallback: float) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return fallback


def _profile_shape() -> dict[str, Any]:
    return {
        "headline": "short professional headline",
        "target_roles": ["primary role suggestions"],
        "secondary_roles": ["adjacent role suggestions"],
        "suggested_roles_reasoning": "brief reason for role suggestions",
        "skills": [
            {
                "name": "skill label",
                "category": "skill category",
                "level": "strong",
                "evidence": "short evidence from CV",
            }
        ],
        "industries": ["industries or domains"],
        "preferred_locations": ["locations if evident"],
        "preferred_work_modes": ["remote | hybrid | onsite if evident"],
        "dealbreakers": ["clear constraints if stated"],
        "avoid_roles": ["roles the profile appears poorly suited for"],
        "real_experience_years": 0,
        "notes": "truthful notes for job ranking",
        "extraction_notes": ["ambiguities the user should review manually"],
        "confidence": "high | medium | low",
    }


def _profile_validation_error(profile: dict[str, Any]) -> str | None:
    problems = []
    roles = [*profile.get("target_roles", []), *profile.get("secondary_roles", [])]
    skills = profile.get("skills") or []
    if not profile.get("headline"):
        problems.append("headline is required")
    if not roles:
        problems.append("at least one target_roles or secondary_roles value is required")
    if not skills:
        problems.append("at least one skill with evidence is required")
    missing_evidence = [skill.get("name") for skill in skills if not skill.get("evidence")]
    if missing_evidence:
        problems.append(f"skills missing evidence: {missing_evidence[:10]}")
    placeholder_levels = [
        skill.get("name")
        for skill in skills
        if "|" in str(skill.get("level") or "") or str(skill.get("level") or "").lower() not in {"strong", "medium", "weak"}
    ]
    if placeholder_levels:
        problems.append(f"skills with invalid level values: {placeholder_levels[:10]}")
    return "; ".join(problems) if problems else None


def _confidence(value: Any) -> str:
    confidence = str(value or "medium").strip().lower()
    return confidence if confidence in {"high", "medium", "low"} else "medium"
